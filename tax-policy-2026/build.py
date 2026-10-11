"""由 sources/ 下的原始检索数据生成附录，并将 Markdown 导出为 Word。

用法：python3 build.py   （依赖 python-docx）
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).parent
SRC = ROOT / "sources"
MAIN = ROOT / "2026年1-9月涉税政策汇编宣讲材料.md"


def build_index():
    docs = json.loads((SRC / "sta_fgk_2026.json").read_text(encoding="utf-8"))
    docs = [d for d in docs if "2026-01-01" <= d["date"] <= "2026-09-30"]
    lines = [
        "# 附录一 2026年1—9月国家税务总局政策法规库文件索引",
        "",
        f"> 来源：国家税务总局政策法规库“最新文件”栏目，共{len(docs)}件，按发布日期倒序排列。",
        "",
        "| 序号 | 发布日期 | 文号 | 标题 |",
        "|---|---|---|---|",
    ]
    for i, d in enumerate(docs, 1):
        title = d["title"].replace("|", "｜")
        lines.append(f"| {i} | {d['date']} | {d['no'] or '—'} | [{title}]({d['url']}) |")
    return "\n".join(lines) + "\n"


def build_qa():
    text = (SRC / "zj_rdwd_2026.txt").read_text(encoding="utf-8")
    out = [
        "# 附录二 浙江省税务局“热点问答”2026年原文汇编",
        "",
        "> 来源：https://zhejiang.chinatax.gov.cn/col/col13314/index.html",
        "",
    ]
    month = None
    for block in re.split(r"\n(?=## \[)", text.strip()):
        head, _, body = block.partition("\n")
        m = re.match(r"## \[(\d{4})-(\d{2})-(\d{2})\] (.*)", head)
        if not m:
            continue
        if m.group(2) != month:
            month = m.group(2)
            out += [f"## 2026年{int(month)}月", ""]
        out += [f"### {m.group(4)}（{m.group(1)}-{m.group(2)}-{m.group(3)}）", ""]
        for line in body.splitlines():
            line = line.strip()
            if not line or line.startswith("问："):
                continue
            out.append("**答：**" if line == "答：" else line)
            out.append("")
    return "\n".join(out)


def build_replies():
    body = (SRC / "zj_zxdf_2026.md").read_text(encoding="utf-8")
    head = (
        "# 附录三 浙江省税务局“咨询答复”2026年原文汇编\n\n"
        "> 来源：https://zhejiang.chinatax.gov.cn/col/col25980/index.html ，答复单位：浙江财税12366服务中心\n\n"
    )
    body = body.replace("\n**答：**", "\n**答：**\n\n")
    return head + re.sub(r"(?<=[^\n])\n(?=[^\n])", "\n\n", body)


def build_cases():
    text = (SRC / "zj_cases_2026.txt").read_text(encoding="utf-8")
    blocks = sorted(re.split(r"\n(?=## )", text.strip()), key=lambda b: b[3:13])
    out = [
        "# 附录四 浙江省税务部门2026年曝光案件原文汇编",
        "",
        "> 来源：国家税务总局浙江省税务局“税收新闻”栏目及各市税务局发布",
        "",
    ]
    for b in blocks:
        head, _, body = b.partition("\n")
        out += [head.replace("## ", "## ", 1), ""]
        paras = []
        for line in (l.strip() for l in body.splitlines()):
            if not line:
                continue
            if paras and not paras[-1].endswith(("。", "：", "”", "！", "？")):
                paras[-1] += line
            else:
                paras.append(line)
        out += [p + "\n" for p in paras]
    return "\n".join(out)


def set_font(run, size=None, bold=None):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


INLINE = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text, size=10.5):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), size, True)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            set_font(par.add_run(m.group(1)), size)
        else:
            set_font(par.add_run(part), size)


def md_to_docx(md, path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[-| :]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-| :]+\|$", lines[i].strip()):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row[: len(rows[0])]):
                    p = table.cell(r, c).paragraphs[0]
                    add_inline(p, cell.replace("<br>", "\n"), 9)
                    if r == 0:
                        for run in p.runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,4}) (.*)", line)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=level)
            run = h.add_run(m.group(2))
            set_font(run, {1: 18, 2: 15, 3: 13, 4: 12}[level], True)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("> "):
            p = doc.add_paragraph()
            add_inline(p, line[2:], 9.5)
            for run in p.runs:
                run.italic = True
        elif re.match(r"^\s*[-*] ", line):
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
            add_inline(p, line.lstrip()[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            add_inline(p, line)
        elif line.strip() == "---":
            pass
        elif line.strip():
            add_inline(doc.add_paragraph(), line)
        i += 1
    doc.save(path)


def main():
    appendices = {
        "附录一_政策文件索引.md": build_index(),
        "附录二_浙江热点问答原文.md": build_qa(),
        "附录三_浙江咨询答复原文.md": build_replies(),
        "附录四_浙江曝光案件原文.md": build_cases(),
    }
    for name, content in appendices.items():
        (ROOT / name).write_text(content, encoding="utf-8")
    md_to_docx(MAIN.read_text(encoding="utf-8"), ROOT / MAIN.with_suffix(".docx").name)
    combined = "\n\n".join(appendices.values())
    md_to_docx(combined, ROOT / "附录_原文汇编.docx")


if __name__ == "__main__":
    main()
